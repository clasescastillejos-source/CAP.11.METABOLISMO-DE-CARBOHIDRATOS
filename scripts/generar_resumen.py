#!/usr/bin/env python3
"""Genera la guía de estudio en DOCX desde el capítulo y contenido editorial.

Las tablas y los cuadros sinópticos son tablas nativas de Word. Las figuras F
son recortes del PDF proporcionado; los esquemas Q son dibujos de apoyo propios.
Todas las fuentes editables, incluidos títulos, tablas y pies, son de 14 puntos.
"""
from __future__ import annotations
import re
import json
import math
from pathlib import Path
from io import BytesIO
import pymupdf
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'CAP.11.METABOLISMO DE CARBOHIDRATOS_compressed.pdf'
CONTENT = ROOT / 'scripts/resumen_contenido.md'
CACHE = ROOT / '.cache/guia'
OUT = ROOT / 'entregables/Resumen_detallado_Metabolismo_de_carbohidratos.docx'
NAVY = '19394E'
TEAL = '007A82'
MAGENTA = 'A51C61'
PALE = 'EAF3F5'
INK = '243746'
GRAY = '596875'
FONT = 'Arial'
SIZE = 14

# Coordenadas en puntos PDF. La numeración PDF es 1-based, libro = PDF + 242.
FIGURES = {
 'fig01': (2, (44, 293, 595, 808), '11-1', 'Panorama de las rutas de la glucosa', '1.3'),
 'fig02': (5, (44, 63, 596, 411), '11-2', 'Digestión del almidón y absorción intestinal', '2.3'),
 'fig03': (8, (44, 70, 594, 619), '11-3', 'Las diez reacciones y las dos fases de la glucólisis', '3.3'),
 'fig04': (10, (44, 67, 592, 524), '11-4', 'Regulación alostérica y hormonal de la glucólisis', '5.3'),
 'fig05': (12, (44, 405, 592, 731), '11-5', 'Entrada de otros monosacáridos en la glucólisis', '6.1'),
 'fig06': (14, (222, 396, 592, 789), '11-6', 'Glucólisis frente a gluconeogénesis', '7.2'),
 'fig07': (16, (221, 425, 592, 800), '11-7', 'Regulación recíproca y ciclos de sustrato', '7.4'),
 'fig08': (17, (43, 432, 414, 802), '11-8', 'Ciclo de Cori y ciclo de la glucosa-alanina', '8.1'),
 'fig09': (18, (224, 555, 594, 802), '11-9', 'Destinos del piruvato y regeneración de NAD⁺', '9.1'),
 'fig10': (19, (44, 686, 412, 807), '11-10', 'Descarboxilación oxidativa del piruvato', '9.4'),
 'fig11': (20, (224, 646, 594, 800), '11-11', 'Regulación del complejo piruvato deshidrogenasa', '9.6'),
 'fig12': (21, (44, 287, 416, 800), '11-12', 'Fases oxidativa y no oxidativa de las pentosas fosfato', '10.2'),
 'fig13': (24, (223, 66, 593, 497), '11-13', 'Activación de la glucosa como UDP-glucosa', '11.2'),
 'fig14': (25, (44, 75, 416, 282), '11-14', 'Mecanismo de la enzima ramificante', '11.3'),
 'fig15': (25, (44, 607, 416, 806), '11-15', 'Mecanismo de la enzima desramificante', '11.4'),
 'fig16': (26, (223, 530, 636, 735), '11-16', 'Control hormonal del metabolismo del glucógeno', '11.6'),
 'gk': (11, (52, 340, 592, 802), 'R11-3', 'Glucoquinasa: compartimentación y secreción de insulina', '5.4'),
 'integracion_hormonal': (28, (53, 103, 585, 463), 'R270-A', 'Esquema original de regulación hormonal', '12.4'),
 'integracion_alosterica': (28, (53, 480, 585, 796), 'R270-B', 'Esquema original de regulación alostérica', '12.4'),
}

# Proyecciones lineales/Fischer: OH de los centros de la serie D a la derecha
# salvo C3 de las hexosas. El símbolo P representa PO3(2-), unido por oxígeno.
MOLECULES = {
 'glucosa': ('D-glucosa', 'CHO', [('H','OH'),('OH','H'),('H','OH'),('H','OH')], 'CH₂OH'),
 'g6p': ('Glucosa-6-fosfato', 'CHO', [('H','OH'),('OH','H'),('H','OH'),('H','OH')], 'CH₂–O–P'),
 'f6p': ('Fructosa-6-fosfato', 'CH₂OH', [('keto','O'),('OH','H'),('H','OH'),('H','OH')], 'CH₂–O–P'),
 'f16': ('Fructosa-1,6-bisfosfato', 'CH₂–O–P', [('keto','O'),('OH','H'),('H','OH'),('H','OH')], 'CH₂–O–P'),
 'dhap': ('Dihidroxiacetona fosfato', 'CH₂OH', [('keto','O')], 'CH₂–O–P'),
 'gap': ('D-gliceraldehído-3-fosfato', 'CHO', [('H','OH')], 'CH₂–O–P'),
 'bpg13': ('1,3-bisfosfoglicerato', 'P–O–C(=O)', [('H','OH')], 'CH₂–O–P'),
 'pg3': ('3-fosfoglicerato', 'COO⁻', [('H','OH')], 'CH₂–O–P'),
 'pg2': ('2-fosfoglicerato', 'COO⁻', [('H','O–P')], 'CH₂OH'),
 'pep': ('Fosfoenolpiruvato', 'COO⁻', [('ene','O–P')], 'CH₂'),
 'pyr': ('Piruvato', 'COO⁻', [('keto','O')], 'CH₃'),
}
REACTIONS = {
 1: ('glucosa','g6p'), 2: ('g6p','f6p'), 3: ('f6p','f16'),
 4: ('f16','gap','dhap'), 5: ('dhap','gap'), 6: ('gap','bpg13'),
 7: ('bpg13','pg3'), 8: ('pg3','pg2'), 9: ('pg2','pep'), 10: ('pep','pyr'),
}


def shade(cell, color):
    pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:fill'), color); pr.append(el)


def cell_margins(cell, size=100):
    pr = cell._tc.get_or_add_tcPr(); mar = OxmlElement('w:tcMar')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + edge); el.set(qn('w:w'), str(size)); el.set(qn('w:type'), 'dxa'); mar.append(el)
    pr.append(mar)


def set_run(run, bold=None, color=None):
    run.font.name = FONT; run.font.size = Pt(SIZE)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)
    return run


def hyperlink(p, label, url=None, anchor=None):
    h = OxmlElement('w:hyperlink')
    if url: h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    if anchor: h.set(qn('w:anchor'), anchor)
    rr = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    ff = OxmlElement('w:rFonts'); ff.set(qn('w:ascii'), FONT); ff.set(qn('w:hAnsi'), FONT); props.append(ff)
    ss = OxmlElement('w:sz'); ss.set(qn('w:val'), '28'); props.append(ss)
    color = OxmlElement('w:color'); color.set(qn('w:val'), TEAL); props.append(color)
    und = OxmlElement('w:u'); und.set(qn('w:val'), 'single'); props.append(und)
    rr.append(props); t = OxmlElement('w:t'); t.text = label; rr.append(t); h.append(rr); p._p.append(h)


def inline(p, text, bold_default=False):
    pattern = r'(\*\*.*?\*\*|\[[^\]]+\]\((?:[^()]|\([^()]*\))*\)|<br>)'
    for part in re.split(pattern, text):
        if not part: continue
        if part == '<br>': set_run(p.add_run()).add_break(); continue
        if part.startswith('**') and part.endswith('**'): set_run(p.add_run(part[2:-2]), True); continue
        m = re.fullmatch(r'\[([^\]]+)\]\(((?:[^()]|\([^()]*\))*)\)', part)
        if m: hyperlink(p, m.group(1), m.group(2)); continue
        set_run(p.add_run(part), bold_default)


def bookmark(p, name, num):
    a=OxmlElement('w:bookmarkStart'); a.set(qn('w:id'),str(num)); a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'),str(num))
    p._p.insert(0,a); p._p.append(b)


def field(p, instruction):
    rr=set_run(p.add_run()); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); rr._r.append(begin)
    rr=set_run(p.add_run()); tx=OxmlElement('w:instrText'); tx.set(qn('xml:space'),'preserve'); tx.text=' '+instruction+' '; rr._r.append(tx)
    rr=set_run(p.add_run()); sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); rr._r.append(sep)
    set_run(p.add_run('1'))
    rr=set_run(p.add_run()); end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); rr._r.append(end)


def extract_figures():
    CACHE.mkdir(parents=True, exist_ok=True)
    pdf=pymupdf.open(SOURCE)
    for key,(pg,box,no,title,loc) in FIGURES.items():
        dst=CACHE/f'{key}.png'
        pix=pdf[pg-1].get_pixmap(dpi=300, clip=pymupdf.Rect(box), alpha=False)
        im=Image.open(BytesIO(pix.tobytes('png')))
        im.save(dst, optimize=True, dpi=(300,300))
    pdf.close()


def text_center(draw, xy, text, font, fill=INK):
    draw.text(xy,text,font=font,fill='#'+fill,anchor='mm')


def wrap_name(text, limit=23):
    if len(text)<=limit: return text
    candidates=[m.start()+1 for m in re.finditer('-',text)] + [m.start() for m in re.finditer(' ',text)]
    if not candidates: return text
    point=min(candidates,key=lambda x:abs(x-len(text)/2))
    return text[:point].strip()+'\n'+text[point:].strip()


def draw_structure(draw, key, x, top, bottom, fonts):
    title,head,rows,tail=MOLECULES[key]
    f,big,small=fonts
    n=len(rows); step=132
    height=(n+1)*step
    yy=(top+bottom-height)/2
    centers=[yy+i*step for i in range(n+2)]
    # Se representan enlaces verticales y grupos; no son modelos 3D.
    for i in range(n+1):
        y1=centers[i]+36; y2=centers[i+1]-36
        draw.line((x,y1,x,y2),fill='#'+INK,width=7)
        if i==n and rows and rows[-1][0]=='ene':
            draw.line((x+20,y1,x+20,y2),fill='#'+INK,width=7)
    if key == 'bpg13':
        y=centers[0]
        text_center(draw,(x,y),'C',f)
        for off in (-12,12): draw.line((x-120,y+off,x-40,y+off),fill='#'+INK,width=7)
        text_center(draw,(x-182,y),'O',f)
        draw.line((x+40,y,x+120,y),fill='#'+INK,width=7)
        text_center(draw,(x+235,y),'O–P',f,MAGENTA)
    else:
        text_center(draw,(x,centers[0]),head,f)
    for i,(left,right) in enumerate(rows):
        y=centers[i+1]
        text_center(draw,(x,y),'C',f)
        if left in ('keto','ene'):
            for offset in (-12,12) if left=='keto' else (0,):
                draw.line((x+40,y+offset,x+120,y+offset),fill='#'+INK,width=7)
            text_center(draw,(x+182,y),right,f, MAGENTA if 'P' in right else INK)
        else:
            draw.line((x-115,y,x-38,y),fill='#'+INK,width=7)
            draw.line((x+38,y,x+115,y),fill='#'+INK,width=7)
            text_center(draw,(x-178,y),left,f, MAGENTA if 'P' in left else INK)
            text_center(draw,(x+178,y),right,f, MAGENTA if 'P' in right else INK)
    text_center(draw,(x,centers[-1]),tail,f, MAGENTA if 'P' in tail else INK)
    draw.multiline_text((x,114),wrap_name(title),font=big,fill='#'+NAVY,anchor='mm',align='center',spacing=8)


def make_reactions():
    regular=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',65)
    bold=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',57)
    small=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',40)
    for num,keys in REACTIONS.items():
        H=(max(len(MOLECULES[k][2]) for k in keys)+1)*132+350
        im=Image.new('RGB',(2400,H),'white'); d=ImageDraw.Draw(im)
        d.rounded_rectangle((15,15,2385,H-15),radius=32,outline='#D0DEE6',width=4)
        if len(keys)==2:
            positions=[590,1805]; Y=(220+H-130)/2; arrow=(1070,Y,1320,Y)
        else:
            positions=[470,1370,2090]; Y=(220+H-130)/2; arrow=(835,Y,1015,Y)
            text_center(d,(1715,Y),'+',regular,TEAL)
        for key,x in zip(keys,positions): draw_structure(d,key,x,220,H-130,(regular,bold,small))
        x1,y1,x2,y2=arrow
        d.line((x1,y1,x2,y2),fill='#'+TEAL,width=12)
        d.polygon([(x2,y2),(x2-38,y2-25),(x2-38,y2+25)],fill='#'+TEAL)
        if num not in (1,3,10):
            d.line((x2,y2+52,x1,y1+52),fill='#'+TEAL,width=9)
            d.polygon([(x1,y1+52),(x1+32,y1+32),(x1+32,y1+72)],fill='#'+TEAL)
        text_center(d,(1200,H-47),'P = PO₃²⁻ · Conectividad simplificada; proyección lineal de la serie D.',small,GRAY)
        im.save(CACHE/f'reaccion_{num:02d}.png',optimize=True,dpi=(350,350))


def new_document():
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8)
    sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
    sec.header_distance=Cm(0.7); sec.footer_distance=Cm(0.7)
    sec.different_first_page_header_footer=True
    for style in doc.styles:
        if style.type in (1,2):
            style.font.name=FONT; style.font.size=Pt(SIZE)
    normal=doc.styles['Normal']
    normal.font.color.rgb=RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing=1.08
    normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.widow_control=True
    for name in ['Heading 1','Heading 2','Heading 3','Title','Subtitle']:
        style=doc.styles[name]; style.font.name=FONT; style.font.size=Pt(SIZE)
        style.font.color.rgb=RGBColor.from_string(NAVY if name in ('Heading 1','Title') else TEAL)
        style.font.bold=True
        style.paragraph_format.space_before=Pt(12)
        style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.keep_with_next=True
        style.paragraph_format.keep_together=True
    doc.styles['Heading 1'].paragraph_format.page_break_before=True
    doc.styles['Caption'].font.name=FONT; doc.styles['Caption'].font.size=Pt(SIZE)
    doc.styles['Caption'].font.color.rgb=RGBColor.from_string(GRAY)
    doc.styles['Caption'].paragraph_format.line_spacing=1.05
    doc.styles['Caption'].paragraph_format.space_after=Pt(8)
    doc.styles['List Bullet'].font.size=Pt(SIZE)
    doc.styles['List Bullet'].paragraph_format.line_spacing=1.08
    doc.styles['List Bullet'].paragraph_format.space_after=Pt(5)
    hdr=sec.header.paragraphs[0]; hdr.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_run(hdr.add_run('CAPÍTULO 11 · GUÍA DE ESTUDIO'), True, TEAL)
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_run(f.add_run('Metabolismo de carbohidratos · '),False,GRAY); field(f,'PAGE')
    settings=doc.settings.element
    upd=OxmlElement('w:updateFields'); upd.set(qn('w:val'),'true'); settings.append(upd)
    no_compress=OxmlElement('w:doNotAutoCompressPictures'); no_compress.set(qn('w:val'),'true'); settings.append(no_compress)
    lang=OxmlElement('w:lang'); lang.set(qn('w:val'),'es-MX')
    doc.styles['Normal'].element.get_or_add_rPr().append(lang)
    doc.core_properties.title='Metabolismo de carbohidratos: resumen detallado y guía de estudio'
    doc.core_properties.subject='Capítulo 11. Glucólisis, gluconeogénesis, piruvato, pentosas fosfato y glucógeno'
    doc.core_properties.author='Material didáctico elaborado a partir del capítulo proporcionado'
    doc.core_properties.keywords='glucólisis, enzimas, estructuras, coenzimas, cofactores, cuadros editables'
    return doc


def table(doc, rows, widths=None, synopsis=False):
    n=len(rows[0]); t=doc.add_table(rows=0, cols=n); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths is None:
        widths=[17.4/n]*n
        if n==2: widths=[5.4,12.0]
        if n==3: widths=[4.0,6.5,6.9]
        if n==4: widths=[2.0,5.5,5.7,4.2]
        if n==5: widths=[1.4,6.6,3.2,3.2,3.0]
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    for ridx,data in enumerate(rows):
        row=t.add_row(); props=row._tr.get_or_add_trPr()
        nobreak=OxmlElement('w:cantSplit'); props.append(nobreak)
        if ridx==0:
            rep=OxmlElement('w:tblHeader'); props.append(rep)
        for cidx,(cell,txt) in enumerate(zip(row.cells,data)):
            cell.width=Cm(widths[cidx]); cell_margins(cell)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]; p.paragraph_format.line_spacing=1.02
            p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
            inline(p,txt,bold_default=(ridx==0))
            if ridx==0:
                shade(cell,NAVY)
                for run in p.runs: set_run(run,True,'FFFFFF')
            elif synopsis and cidx==0:
                shade(cell,PALE)
                for run in p.runs: set_run(run,True,TEAL)
            elif ridx%2==0: shade(cell,'F4F7FA')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t


def note(doc,title,text):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.columns[0].width=Cm(17.4)
    c=t.cell(0,0); cell_margins(c,140); shade(c,PALE)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(5); set_run(p.add_run(title),True,TEAL)
    p=c.add_paragraph(); p.paragraph_format.space_after=Pt(3); inline(p,text)
    no=OxmlElement('w:cantSplit'); t.rows[0]._tr.get_or_add_trPr().append(no)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)


def image(doc,path,caption,alt,max_height=14.2,width=17.4,bm=None,bm_id=500):
    with Image.open(path) as im: w,h=im.size
    actual_width=min(width,max_height*w/h)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4); p.paragraph_format.keep_with_next=True
    if bm: bookmark(p,bm,bm_id)
    pic=p.add_run().add_picture(str(path),width=Cm(actual_width)); pic._inline.docPr.set('descr',alt)
    cap=doc.add_paragraph(style='Caption'); inline(cap,caption)
    return p


def cover(doc):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(38)
    set_run(p.add_run('BIOQUÍMICA · CAPÍTULO 11'),True,TEAL)
    p=doc.add_paragraph(style='Title'); p.paragraph_format.space_before=Pt(24)
    set_run(p.add_run('METABOLISMO DE LOS HIDRATOS DE CARBONO'),True,NAVY)
    p=doc.add_paragraph(style='Subtitle'); p.paragraph_format.space_before=Pt(14)
    set_run(p.add_run('Resumen extenso, explicado y orientado al aprendizaje'),True,TEAL)
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(15)
    inline(p,'**Énfasis especial en la glucólisis:** diez reacciones, estructuras, enzimas, coenzimas, cofactores, regulación y balance energético.')
    t=table(doc,[['IDEA CENTRAL','RESULTADO DE LA GLUCÓLISIS'],['1 glucosa de 6 carbonos','→ 2 piruvatos de 3 carbonos'],['Balance neto de la vía','2 ATP + 2 NADH por glucosa']],widths=[6.0,11.4],synopsis=True)
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(18)
    inline(p,'Incluye las 16 figuras numeradas del libro, tres láminas complementarias del capítulo y diez esquemas químicos propios de alta resolución. Los cuadros comparativos y sinópticos son editables en Word.')
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(20)
    inline(p,'**Fuente principal:** archivo «CAP.11.METABOLISMO DE CARBOHIDRATOS_compressed.pdf», páginas 243–270 del libro (28 páginas del PDF).')
    p=doc.add_paragraph(); inline(p,'**Formato:** A4 · texto, títulos, cuadros y pies de 14 puntos · elaboración: 30 de septiembre de 2026.')
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(14)
    inline(p,'Material de estudio. No sustituye al capítulo original ni a la orientación docente o clínica.')
    doc.add_page_break()


def create():
    extract_figures(); make_reactions(); doc=new_document(); cover(doc)
    raw=CONTENT.read_text(encoding='utf-8').splitlines()
    used=[]; counters={'h':1,'fig':500}; i=0
    while i<len(raw):
        line=raw[i].strip()
        if not line: i+=1; continue
        if line=='@@SALTO': doc.add_page_break(); i+=1; continue
        if line=='@@INDICE':
            p=doc.add_paragraph(style='Heading 2'); inline(p,'Índice de secciones')
            for ln in raw:
                if ln.startswith('# '):
                    title=ln[2:].strip(); n=re.match(r'(\d+)\.',title)
                    if n:
                        pp=doc.add_paragraph(); hyperlink(pp,title,anchor='sec_'+n.group(1))
            i+=1; continue
        if line=='@@INDICEFIGURAS':
            rows=[['Imagen original','Página del libro / PDF','Ubicación en esta guía']]
            for k in used:
                pg,box,num,title,loc=FIGURES[k]
                rows.append([f'{num}: {title}',f'{pg+242} / {pg}',f'§ {loc}'])
            table(doc,rows,widths=[8.0,3.3,6.1]); i+=1; continue
        if line.startswith('@@FIG '):
            key=line.split()[1]; pg,box,num,title,loc=FIGURES[key]
            prefix='Figura '+num if num.startswith('11-') else ('Lámina del recuadro 11-3' if key=='gk' else 'Lámina de integración '+('A' if key.endswith('hormonal') else 'B'))
            cap=f'**{prefix}. {title}.** Fuente: capítulo proporcionado, p. {pg+242} del libro (p. {pg} del PDF). Ubicación en esta guía: § {loc}. Imagen original recortada, no redibujada.'
            image(doc,CACHE/f'{key}.png',cap,title,max_height=14.6,bm=key,bm_id=counters['fig'])
            counters['fig']+=1; used.append(key); i+=1; continue
        if line.startswith('@@REACCION '):
            num=int(line.split()[1]); title=f'Esquema químico Q{num:02d}: reacción {num} de la glucólisis'
            image(doc,CACHE/f'reaccion_{num:02d}.png',f'**{title}.** Elaboración didáctica propia. P representa PO₃²⁻; el enlace al esqueleto carbonado se realiza a través de oxígeno. Las formas lineales simplifican la comparación estructural.',title,max_height=7.8)
            i+=1; continue
        if line.startswith('@@NOTA '):
            title,text=line[len('@@NOTA '):].split('|',1); note(doc,title.strip(),text.strip()); i+=1; continue
        if line.startswith('@@ECUACION '):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(9)
            inline(p,line[len('@@ECUACION '):],True)
            for r in p.runs: r.font.color.rgb=RGBColor.from_string(TEAL)
            i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(raw) and raw[i].strip().startswith('|'):
                vals=[v.strip() for v in raw[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',v or '-') for v in vals): rows.append(vals)
                i+=1
            if rows:
                assert all(len(r)==len(rows[0]) for r in rows),rows
                synopsis=any('sinóptico' in v.lower() for v in rows[0])
                table(doc,rows,synopsis=synopsis)
            continue
        m=re.match(r'^(#{1,3})\s+(.*)$',line)
        if m:
            level=len(m.group(1)); text=m.group(2); p=doc.add_paragraph(style=f'Heading {level}'); inline(p,text)
            if level==1:
                n=re.match(r'(\d+)\.',text)
                if n: bookmark(p,'sec_'+n.group(1),counters['h']); counters['h']+=1
            i+=1; continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); inline(p,line[2:]); i+=1; continue
        # Cada línea editorial es un párrafo; no se convierten saltos suaves en párrafos vacíos.
        p=doc.add_paragraph(); inline(p,line); i+=1
    # Fuente explícita de 14 pt para todos los runs, incluso si han heredado otro estilo.
    for p in doc.paragraphs:
        for r in p.runs: r.font.name=FONT; r.font.size=Pt(SIZE)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs: r.font.name=FONT; r.font.size=Pt(SIZE)
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)
    metadata={'archivo':str(OUT.relative_to(ROOT)),'fuente':SOURCE.name,'paginas_fuente':28,
              'letra_puntos':14,'figuras_originales':len(used),'esquemas_quimicos_propios':10,
              'tablas_nativas_editables':len(doc.tables),'palabras_contenido':len(re.findall(r'\b\w+\b',CONTENT.read_text(encoding='utf-8')))}
    (CACHE/'estadisticas.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(metadata,ensure_ascii=False,indent=2))

if __name__=='__main__': create()
