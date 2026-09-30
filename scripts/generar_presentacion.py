#!/usr/bin/env python3
"""Transposición íntegra y verificable del DOCX a PowerPoint y PDF 16:9.

No resume ni reescribe párrafos. Pagina texto y celdas a 20–22 pt; todos los
textos visibles de PowerPoint, incluidos pies y leyendas, tienen >=18 pt.
Las 19 imágenes originales se conservan; las 10 estructuras se reconstruyen
como texto y enlaces vectoriales editables a partir del mismo modelo químico.
"""
from __future__ import annotations
import hashlib
import importlib.util
import json
import re
import math
from pathlib import Path
from collections import Counter, defaultdict
from zipfile import ZipFile
from lxml import etree
from PIL import Image
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.ns import qn as wqn
from pptx import Presentation
from pptx.util import Pt
from presentacion_diseno import Deck, wrap, rich_slice, spans_text, text_height, plain, clean, rich_width
from presentacion_diseno import SW, SH, LEFT, RIGHT, NAVY, INK, TEAL, WHITE, GRAY, LIGHT, LINE, MAGENTA, GOLD, COLORS

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'entregables/Resumen_detallado_Metabolismo_de_carbohidratos.docx'
CACHE=ROOT/'.cache/presentacion'
ASSETS=CACHE/'assets'
OUT=ROOT/'entregables/Presentacion_completa_Metabolismo_de_carbohidratos.pptx'
PDF=ROOT/'entregables/Presentacion_completa_Metabolismo_de_carbohidratos.pdf'
REPORT=ROOT/'entregables/Verificacion_presentacion.json'
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}


def sha(data):return hashlib.sha256(data).hexdigest()


def bold_value(r):
    el=r.find(wqn('w:rPr'))
    if el is None:return False
    b=el.find(wqn('w:b'))
    return b is not None and b.get(wqn('w:val'),'true') not in ('0','false','off')


def para_spans(p):
    spans=[]
    for el in p._p:
        if el.tag==wqn('w:r'):
            bold=bold_value(el)
            for child in el:
                if child.tag==wqn('w:t') and child.text:spans.append({'text':child.text,'bold':bold})
                elif child.tag==wqn('w:br') and child.get(wqn('w:type'))!='page':spans.append({'text':'\n','bold':bold})
                elif child.tag==wqn('w:tab'):spans.append({'text':'\t','bold':bold})
        elif el.tag==wqn('w:hyperlink'):
            rid=el.get(wqn('r:id'));anchor=el.get(wqn('w:anchor'))
            link=p.part.rels[rid].target_ref if rid else None
            for rr in el.findall(wqn('w:r')):
                tt=''.join(rr.itertext()) if False else ''.join(t.text or '' for t in rr.iter(wqn('w:t')))
                if tt:
                    s={'text':tt,'bold':bold_value(rr)}
                    if link:s['link']=link
                    if anchor:s['anchor']=anchor
                    spans.append(s)
    return spans


def extract():
    ASSETS.mkdir(parents=True,exist_ok=True)
    doc=Document(SOURCE);blocks=[];units={};pi=0;ti=0;ii=0
    def unit(uid,spans,kind,**extra):
        value={'id':uid,'spans':spans,'text':spans_text(spans),'kind':kind,**extra};units[uid]=value;return value
    for el in doc.element.body:
        if el.tag==wqn('w:p'):
            p=Paragraph(el,doc);ss=para_spans(p);tt=spans_text(ss)
            if tt.strip():
                pi+=1;u=unit(f'P{pi:04d}',ss,'paragraph',style=p.style.name)
                blocks.append({'kind':'paragraph','unit':u})
            for pic in el.xpath('.//wp:inline'):
                ii+=1;rid=pic.xpath('.//a:blip')[0].get(wqn('r:embed'))
                part=doc.part.related_parts[rid];blob=part.blob;ext=part.content_type.split('/')[-1]
                if ext=='jpeg':ext='jpg'
                dst=ASSETS/f'original_{ii:02d}.{ext}';dst.write_bytes(blob)
                descr=pic.xpath('wp:docPr')[0].get('descr','')
                blocks.append({'kind':'image','id':f'I{ii:02d}','path':str(dst),'sha256':sha(blob),'alt':descr})
        elif el.tag==wqn('w:tbl'):
            ti+=1;rows=[]
            for r,row in enumerate(Table(el,doc).rows):
                data=[]
                for c,cell in enumerate(row.cells):
                    ss=[]
                    for k,p in enumerate(cell.paragraphs):
                        if k:ss.append({'text':'\n','bold':False})
                        ss.extend(para_spans(p))
                    u=unit(f'T{ti:03d}R{r:03d}C{c:02d}',ss,'cell',table=ti,row=r,col=c)
                    data.append(u)
                rows.append(data)
            blocks.append({'kind':'table','id':f'T{ti:03d}','rows':rows})
    # Asociar cada imagen con su pie exacto, sin perder ese párrafo.
    out=[];i=0
    while i<len(blocks):
        b=blocks[i]
        if b['kind']=='image' and i+1<len(blocks) and blocks[i+1]['kind']=='paragraph' and blocks[i+1]['unit']['style']=='Caption':
            b['caption']=blocks[i+1]['unit'];i+=1
            m=re.search(r'Esquema químico Q(\d{2})',b['caption']['text'])
            if m:b['reaction']=int(m.group(1))
            else:
                m=re.search(r'Figura (11-\d+)',b['caption']['text'])
                b['figure']=m.group(1) if m else ('R11-3' if 'recuadro 11-3' in b['caption']['text'] else ('R270-A' if 'integración A' in b['caption']['text'] else 'R270-B'))
        out.append(b);i+=1
    data={'source':str(SOURCE.name),'sha256':sha(SOURCE.read_bytes()),'units':units,'blocks':out}
    (CACHE/'fuente.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    return data


def model_chemistry():
    spec=importlib.util.spec_from_file_location('word_model',ROOT/'scripts/generar_resumen.py')
    mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    return mod.MOLECULES,mod.REACTIONS


class Composer:
    def __init__(self,data):
        self.data=data;self.deck=Deck();self.chapter=0;self.title='Datos introductorios del resumen original de Word'
        self.heading=None;self.previous_heading=None;self.header_pending=[];self.buffer=[];self.repeated=Counter();self.figs=[];self.tables=0
        self.molecules,self.reactions=model_chemistry()

    def source(self,u,a=0,b=None):return (u['id'],a,len(u['text']) if b is None else b)

    def cover(self,units):
        d=self.deck;d.new(0,dark=True)
        y=27
        for u,size,color,gap in zip(units,[18,40,24,22],['A6E7E7',WHITE,'5DE2D2','DDEBF0'],[36,30,25,0]):
            ss=[{**s,'bold':True} for s in u['spans']] if size==40 else u['spans']
            ls=wrap(ss,RIGHT-LEFT,size);h=text_height(ls,size,1.10)
            d.text(u['text'],LEFT,y,RIGHT-LEFT,h,size,color,spans=ss,lines=ls,leading=1.10,source=self.source(u));y+=h+gap
        # Motivo visual: conservación de los seis carbonos, no una fórmula 3D.
        y=428
        for i in range(6):
            x=LEFT+i*39;d.circle(x,y,27,'28576C')
            d.text('C',x,y+2,27,23,18,WHITE,align='center')
        d.line(313,y+14,425,y+14,'55DCCB',2.8,True)
        for g in range(2):
            for i in range(3):
                x=472+g*173+i*39;d.circle(x,y,27,TEAL)
                d.text('C',x,y+2,27,23,18,WHITE,align='center')
        d.notes('Portada con los primeros cuatro párrafos completos del resumen de Word. Los grupos de círculos solo representan conservación del número de carbonos.')

    def presentation_note(self):
        d=self.deck;d.new(0,'Cómo está organizada esta presentación')
        self.local_paragraph('Se conserva íntegramente la información del resumen de Word: párrafos, celdas de cuadros, explicaciones de imágenes, balances, referencias, glosario y ejercicios. No se han reemplazado por versiones abreviadas.',d.current_scene()['body_top'],22)
        y=d.current_scene()['body_top']+139
        self.local_paragraph('El texto editable utiliza tamaños de 18 puntos o mayores. La información se distribuye en tantas diapositivas como requiere su lectura. Las indicaciones de formato del Word original se mantienen como datos de ese documento, no como el formato de esta presentación.',y,22)
        y+=137
        self.local_paragraph('Las imágenes originales del libro se conservan y se acompañan de vistas ampliadas. Su nitidez sigue limitada por el PDF comprimido. Los diez esquemas químicos se reconstruyen como elementos editables y legibles.',y,20)
        d.notes('Esta diapositiva añade instrucciones de uso; no sustituye ningún contenido del Word.')

    def local_paragraph(self,text,y,size=22):
        ss=plain(text);ls=wrap(ss,RIGHT-LEFT-28,size);h=text_height(ls,size)+16
        self.deck.rect(LEFT,y,RIGHT-LEFT,h,LIGHT,radius=True)
        self.deck.text(text,LEFT+14,y+8,RIGHT-LEFT-28,h-16,size,INK,spans=ss,lines=ls)
        return h

    def start_content(self,title=None):
        title=title or self.title
        base=title;self.repeated[base]+=1
        label=base if self.repeated[base]==1 else base+' — continuación'
        source=None
        if self.heading and self.heading['id'] in self.header_pending:
            source=self.source(self.heading);self.header_pending.remove(self.heading['id'])
        self.deck.new(self.chapter,label,title_source=source)
        # Un encabezado intermedio anterior que no tuvo cuerpo se conserva en el siguiente panel.
        y=self.deck.current_scene()['body_top']
        for uid in self.header_pending[:]:
            u=self.data['units'][uid];ls=wrap(plain(u['text'],True),RIGHT-LEFT,24);h=text_height(ls,24)
            self.deck.text(u['text'],LEFT,y,RIGHT-LEFT,h,24,TEAL,bold=True,lines=ls,source=self.source(u));y+=h+10
            self.header_pending.remove(uid)
        self.deck.current_scene()['body_top']=y
        return y,self.deck.current_scene()['body_bottom']

    def divider(self,u):
        self.flush()
        if self.header_pending:self.start_content();self.header_pending.clear()
        n=re.match(r'(\d+)\.',u['text']);self.chapter=int(n.group(1)) if n else self.chapter
        d=self.deck;d.new(self.chapter,dark=True)
        d.current_scene()['chapter_divider']=True;d.current_scene()['title']=u['text']
        d.text('CAPÍTULO 11 · SECCIÓN '+str(self.chapter),LEFT,30,RIGHT-LEFT,26,18,'99DEDF',bold=True)
        d.text(f'{self.chapter:02d}',LEFT,80,190,100,80,'37D3C6',bold=True,leading=1.0)
        ss=plain(u['text'],True);ls=wrap(ss,RIGHT-LEFT,38);h=text_height(ls,38,1.14)
        d.text(u['text'],LEFT,215,RIGHT-LEFT,h,38,WHITE,bold=True,lines=ls,leading=1.14,source=self.source(u))
        d.text('Contenido íntegro · lectura didáctica · cuadros y estructuras',LEFT,430,RIGHT-LEFT,30,20,'D1E6ED')
        d.targets['sec_'+str(self.chapter)]=d.current_slide();d.notes('Inicio de sección; se conserva el encabezado completo del documento de Word.')
        self.heading=None;self.title=u['text']

    def new_heading(self,u):
        self.flush()
        self.heading=u;self.title=u['text'];self.header_pending.append(u['id'])

    def paragraph(self,u):self.buffer.append(u)

    def flush(self):
        if not self.buffer:return
        queue=[]
        for u in self.buffer:
            ss=u['spans'];font=22
            # El ancho sigue siendo amplio; no se cambian palabras por viñetas resumidas.
            ls=wrap(ss,RIGHT-LEFT-34,font)
            queue.append({'u':u,'lines':ls,'font':font,'start':0,'end':len(u['text'])})
        self.buffer=[]
        idx=0
        while idx<len(queue):
            y,bottom=self.start_content();scene_ids=[]
            while idx<len(queue):
                item=queue[idx];u=item['u'];ls=item['lines'];size=item['font']
                available=bottom-y;h=text_height(ls,size)+18
                if h>available and scene_ids:break
                if h>available:
                    take=max(1,int((available-21)/(size*1.16)))
                    chosen=ls[:take];remainder=ls[take:]
                    a=item['start'];b=chosen[-1]['end']
                    # Los índices de wrap se refieren siempre al texto original completo.
                    part=rich_slice(u['spans'],a,b);txt=spans_text(part)
                    part_lines=wrap(part,RIGHT-LEFT-34,size);ph=text_height(part_lines,size)+18
                    self.deck.rect(LEFT,y,RIGHT-LEFT,ph,LIGHT,radius=True)
                    self.deck.rect(LEFT,y,4,ph,COLORS.get(self.chapter,TEAL))
                    self.deck.text(txt,LEFT+17,y+9,RIGHT-LEFT-34,ph-18,size,INK,spans=part,lines=part_lines,source=self.source(u,a,b))
                    item['start']=b
                    while item['start']<len(u['text']) and u['text'][item['start']].isspace():item['start']+=1
                    rest=rich_slice(u['spans'],item['start']);rl=wrap(rest,RIGHT-LEFT-34,size)
                    # Recalibrar los índices al espacio del texto completo para la próxima porción.
                    for l in rl:l['start']+=item['start'];l['end']+=item['start']
                    item['lines']=rl;scene_ids.append(u['id']);break
                ss=rich_slice(u['spans'],item['start'],item['end']);txt=spans_text(ss)
                use_lines=wrap(ss,RIGHT-LEFT-34,size)
                self.deck.rect(LEFT,y,RIGHT-LEFT,h,LIGHT,radius=True)
                self.deck.rect(LEFT,y,4,h,COLORS.get(self.chapter,TEAL))
                self.deck.text(txt,LEFT+17,y+9,RIGHT-LEFT-34,h-18,size,INK,spans=ss,lines=use_lines,source=self.source(u,item['start'],item['end']))
                scene_ids.append(u['id']);y+=h+12;idx+=1
            self.deck.notes('Transcripción literal de unidades del Word: '+', '.join(scene_ids)+'. El ajuste es únicamente de paginación y tipografía; no se ha condensado el texto.')

    def widths(self,n):
        total=RIGHT-LEFT
        return [total*x for x in ({2:[.32,.68],3:[.28,.33,.39],4:[.105,.395,.29,.21]}.get(n,[1/n]*n))]

    def cell_info(self,u,header=False,first=False,a=0,b=None):
        b=len(u['text']) if b is None else b
        ss=rich_slice(u['spans'],a,b)
        if header or first:ss=[{**s,'bold':True} for s in ss]
        return {'spans':ss,'source':self.source(u,a,b)}

    def row_height(self,row,widths,size,header=False):
        return max(text_height(wrap([{**s,'bold':True} for s in u['spans']] if header or c==0 else u['spans'],widths[c]-20,size),size,1.14)+14 for c,u in enumerate(row))

    def table(self,b):
        self.flush();self.tables+=1;rows=b['rows'];n=len(rows[0])
        if n==1:
            u=rows[0][0];y,bottom=self.start_content();ss=u['spans'];ls=wrap(ss,RIGHT-LEFT-36,22);h=text_height(ls,22)+28
            if h<=bottom-y:
                self.deck.rect(LEFT,y,RIGHT-LEFT,h,'E6F1F2',line='80BFC2',radius=True)
                self.deck.text(u['text'],LEFT+18,y+14,RIGHT-LEFT-36,h-28,22,INK,spans=ss,lines=ls,source=self.source(u))
                self.deck.notes('Recuadro completo del Word, tabla '+b['id']+'. El texto es editable.');return
            self.buffer=[u];self.flush();return
        widths=self.widths(n)
        if n==3 and rows[0][0]['text']=='Paso' and any('fosfogluconolactona' in row[0]['text'] for row in rows[1:]):
            # El nombre químico debe leerse completo, sin cortes dentro de la palabra.
            widths=[(RIGHT-LEFT)*r for r in (.32,.32,.36)]
        size=20;header=rows[0];hh=self.row_height(header,widths,size,True)
        pos=1;part=0
        while pos<len(rows):
            part+=1;y,bottom=self.start_content();maxh=bottom-y
            batch=[header];heights=[hh];nextpos=pos
            while nextpos<len(rows):
                rh=self.row_height(rows[nextpos],widths,size)
                if sum(heights)+rh>maxh:break
                batch.append(rows[nextpos]);heights.append(rh);nextpos+=1
            if nextpos==pos:
                # Caso de fila muy larga: subdividir su contenido, conservando todas las celdas.
                row=rows[pos];frags=[{'u':u,'start':0} for u in row]
                while any(f['start']<len(f['u']['text']) for f in frags):
                    if part>1 or batch!=[header]:y,bottom=self.start_content();maxh=bottom-y
                    headercells=[self.cell_info(u,True,c==0) for c,u in enumerate(header)]
                    maxlines=max(1,int((maxh-hh-17)/(size*1.14)))
                    contents=[];rh=0
                    for c,f in enumerate(frags):
                        u=f['u'];a=f['start'];ss=rich_slice(u['spans'],a);ls=wrap(ss,widths[c]-20,size)
                        chosen=ls[:maxlines];end=a+(chosen[-1]['end'] if chosen else 0)
                        contents.append(self.cell_info(u,False,c==0,a,end))
                        rh=max(rh,text_height(wrap(contents[-1]['spans'],widths[c]-20,size),size,1.14)+14)
                        f['start']=end
                        while f['start']<len(u['text']) and u['text'][f['start']].isspace():f['start']+=1
                    self.deck.table([headercells,contents],widths,LEFT,y,[hh,rh],size,b['id'],part)
                    self.deck.notes('Fila continuada de '+b['id']+'; todas las palabras y celdas se conservan.');part+=1
                pos+=1;continue
            cooked=[]
            for r,row in enumerate(batch):cooked.append([self.cell_info(u,r==0,c==0) for c,u in enumerate(row)])
            self.deck.table(cooked,widths,LEFT,y,heights,size,b['id'],part)
            self.deck.notes('Cuadro nativo y editable de Word '+b['id']+f'. Filas originales {pos}–{nextpos-1}; encabezados repetidos para comprensión. No se elimina ninguna celda.')
            pos=nextpos

    def draw_molecule(self,key,x,top,bottom,name_y,name_w):
        d=self.deck;title,head,rows,tail=self.molecules[key];n=len(rows);step=32;start=(top+bottom-(n+1)*step)/2
        title_lines=wrap(plain(title,True),name_w,23);th=text_height(title_lines,23,1.06)
        d.text(title,x-name_w/2,name_y,name_w,th,23,NAVY,bold=True,align='center',lines=title_lines,leading=1.06)
        yy=[start+k*step for k in range(n+2)]
        for k in range(n+1):
            d.line(x,yy[k]+23,x,yy[k+1]+2,INK,1.8)
            if k==n and rows and rows[-1][0]=='ene':d.line(x+7,yy[k]+23,x+7,yy[k+1]+2,INK,1.8)
        if key=='bpg13':
            y=yy[0];d.text('C',x-15,y,30,31,25,INK,align='center')
            d.text('O',x-68,y,32,31,25,INK,align='center')
            for off in (10,17):d.line(x-41,y+off,x-20,y+off,INK,1.8)
            d.line(x+18,y+14,x+45,y+14,INK,1.8)
            d.text('O–P',x+48,y,76,31,25,MAGENTA,align='center')
        else:d.text(head,x-106,yy[0],212,31,25,INK,align='center')
        for k,(left,right) in enumerate(rows):
            y=yy[k+1];d.text('C',x-15,y,30,31,25,INK,align='center')
            if left in ('keto','ene'):
                for off in ((10,17) if left=='keto' else (14,)):d.line(x+18,y+off,x+45,y+off,INK,1.8)
                d.text(right,x+49,y,83,31,25,MAGENTA if 'P' in right else INK,align='center')
            else:
                d.line(x-45,y+14,x-18,y+14,INK,1.8);d.line(x+18,y+14,x+45,y+14,INK,1.8)
                d.text(left,x-118,y,73,31,25,MAGENTA if 'P' in left else INK,align='center')
                d.text(right,x+49,y,83,31,25,MAGENTA if 'P' in right else INK,align='center')
        d.text(tail,x-106,yy[-1],212,31,25,MAGENTA if 'P' in tail else INK,align='center')

    def chemical(self,b):
        d=self.deck;num=b['reaction'];u=b['caption'];y,bottom=self.start_content()
        # Pie íntegro en una diapositiva propia para no reducir fuentes ni atomizarlo.
        name_y=y+8;top=y+98;bot=bottom-52
        d.rect(LEFT,y,RIGHT-LEFT,bottom-y,'F7FAFC',line=LINE,radius=True)
        keys=self.reactions[num]
        if len(keys)==2:
            xs=[255,709];nw=352
            for key,x in zip(keys,xs):self.draw_molecule(key,x,top,bot,name_y,nw)
            cy=(top+bot)/2+13;d.line(448,cy,517,cy,TEAL,3,True)
            if num not in (1,3,10):d.line(517,cy+21,448,cy+21,TEAL,2.4,True)
        else:
            xs=[190,490,780];nw=260
            for key,x in zip(keys,xs):self.draw_molecule(key,x,top,bot,name_y,nw)
            cy=(top+bot)/2+13;d.line(325,cy,385,cy,TEAL,3,True);d.line(385,cy+21,325,cy+21,TEAL,2.4,True)
            d.text('+',644,cy-9,34,31,25,TEAL,align='center')
        legend='P = PO₃²⁻ · Conectividad simplificada; proyección lineal de la serie D.'
        ls=wrap(plain(legend),RIGHT-LEFT-24,18);h=text_height(ls,18)
        d.text(legend,LEFT+12,bottom-h-4,RIGHT-LEFT-24,h,18,GRAY,lines=ls,align='center')
        d.visuals.append({'source_id':b['id'],'reaction':num,'kind':'native_chemical','slide':len(d.scenes),'model':list(keys)})
        d.notes('Esquema Q%02d redibujado con todos los nombres, grupos químicos y enlaces del modelo del Word. Letras y leyenda >=18 pt; elementos editables. No se ha cambiado la reacción.'%num)
        self.paragraph(u);self.flush()

    def zoom_regions(self,figure):
        # Recortes pedagógicos adicionales. La vista original completa nunca se sustituye.
        return {
          '11-1':[(.27,0,1,.33,'Polímeros de la dieta y glucógeno'),(.27,.31,1,.68,'Glucosa, pentosas y vías opuestas'),(.27,.66,1,1,'Piruvato, lactato y acetil-CoA')],
          '11-2':[(0,0,1,.39,'Panel a: productos de la digestión'),(0,.38,.61,1,'Panel b: enzimas y monosacáridos'),(.57,.38,1,1,'Panel b: transporte intestinal')],
          '11-4':[(.43,0,1,.59,'Panel b: control de fructosa-2,6-bisfosfato'),(.58,.58,1,1,'Panel c: piruvato quinasa hepática')],
          '11-5':[(0,0,.38,1,'Fructosa'),(.37,0,.68,1,'Entrada central a la glucólisis'),(.66,0,1,1,'Galactosa y manosa')],
          '11-6':[(0,0,1,.4,'Rodeos de las fosfatasas'),(0,.64,1,1,'Rodeo piruvato–oxaloacetato–PEP')],
          '11-7':[(0,0,1,.47,'Control de glucosa y fructosa'),(0,.45,1,1,'Control de piruvato y PEP')],
          '11-8':[(0,0,1,.49,'Panel a: ciclo de Cori'),(0,.48,1,1,'Panel b: ciclo glucosa-alanina')],
          '11-12':[(0,0,1,.3,'Panel a: fase oxidativa'),(0,.29,1,.68,'Panel b: pentosas e intercambio de carbono'),(0,.65,1,1,'Panel b: conexión con F6P y GAP')],
          '11-13':[(0,0,1,.42,'Glucosa-6-fosfato y glucosa-1-fosfato'),(0,.40,1,1,'Formación de UDP-glucosa')],
          'R11-3':[(0,0,1,.50,'Panel a: regulación de glucoquinasa'),(.30,.48,1,1,'Panel b: secreción de insulina')],
          'R270-A':[(0,0,1,.54,'Integración hormonal: glucosa, glucógeno y pentosas'),(0,.48,1,1,'Integración hormonal: triosas y piruvato')],
          'R270-B':[(0,0,1,.55,'Integración alostérica: hexosas y reservas'),(0,.45,1,1,'Integración alostérica: PEP y piruvato')],
        }.get(figure,[])

    def image(self,b):
        self.flush()
        if b.get('reaction'):self.chemical(b);return
        u=b['caption'];y,bottom=self.start_content()
        ls=wrap(u['spans'],RIGHT-LEFT-12,18);ch=text_height(ls,18)+9
        ih=bottom-y-ch-15
        if ih<115:
            # Reservar una lámina completa y colocar el pie íntegro después, antes de las ampliaciones.
            self.deck.rect(LEFT,y,RIGHT-LEFT,bottom-y,'F7FAFC',line=LINE,radius=True)
            self.deck.image(b['path'],LEFT+9,y+9,RIGHT-LEFT-18,bottom-y-18,b['alt'],b['id'])
            self.paragraph(u);self.flush()
        else:
            self.deck.rect(LEFT,y,RIGHT-LEFT,ih+12,'F7FAFC',line=LINE,radius=True)
            self.deck.image(b['path'],LEFT+8,y+6,RIGHT-LEFT-16,ih,b['alt'],b['id'])
            self.deck.text(u['text'],LEFT+6,y+ih+21,RIGHT-LEFT-12,ch-9,18,GRAY,spans=u['spans'],lines=ls,source=self.source(u))
        self.figs.append(b);self.deck.visuals.append({'source_id':b['id'],'kind':'original_image','slide':len(self.deck.scenes),'sha256':b['sha256']})
        self.deck.notes('Imagen original íntegra del resumen; SHA256 '+b['sha256']+'. Su pie y su explicación están transcritos sin abreviar. La imagen conserva la nitidez disponible en el PDF comprimido.')
        for k,(a,c,e,f,label) in enumerate(self.zoom_regions(b['figure']),1):
            im=Image.open(b['path']);w,h=im.size
            path=ASSETS/f'ampliacion_{b["id"]}_{k}.png';im.crop((int(a*w),int(c*h),int(e*w),int(f*h))).save(path,optimize=True)
            title=f'{b["figure"]} · Vista ampliada: {label}'
            self.deck.new(self.chapter,title)
            top=self.deck.current_scene()['body_top'];bot=self.deck.current_scene()['body_bottom']
            note='Detalle complementario de la imagen original. La lámina completa y su explicación se conservan en esta sección.'
            nl=wrap(plain(note),RIGHT-LEFT,18);nh=text_height(nl,18)
            self.deck.rect(LEFT,top,RIGHT-LEFT,bot-top-nh-14,'F7FAFC',line=LINE,radius=True)
            self.deck.image(path,LEFT+8,top+7,RIGHT-LEFT-16,bot-top-nh-28,label)
            self.deck.text(note,LEFT,bot-nh,RIGHT-LEFT,nh,18,GRAY,lines=nl)
            self.deck.notes('Recorte adicional para lectura; no reemplaza ni modifica la imagen original. Fuente: '+u['text'])

    def build(self):
        blocks=self.data['blocks'];cover=[];start=0
        while start<len(blocks) and len(cover)<4:
            if blocks[start]['kind']=='paragraph':cover.append(blocks[start]['unit'])
            start+=1
        assert len(cover)==4
        self.cover(cover);self.presentation_note()
        for b in blocks[start:]:
            if b['kind']=='paragraph':
                u=b['unit'];style=u['style']
                if style=='Heading 1':self.divider(u)
                elif style.startswith('Heading') or style in ('Title','Subtitle'):self.new_heading(u)
                else:self.paragraph(u)
            elif b['kind']=='table':self.table(b)
            elif b['kind']=='image':self.image(b)
        self.flush()
        if self.header_pending:self.start_content()
        self.deck.finish();return self.deck


def verify(deck,data):
    prs=Presentation(OUT)
    errors=[];covered=defaultdict(list);allsize=[];cells=0;native_tables=0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text:
                            assert r.font.size is not None,(shape.name,r.text)
                            allsize.append(r.font.size.pt)
            if shape.has_table:
                native_tables+=1
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                if r.text:allsize.append(r.font.size.pt)
    assert allsize and min(allsize)>=18,min(allsize)
    for loc in deck.ledger:
        slide=prs.slides[loc['slide']-1]
        shape=next(s for s in slide.shapes if s.name==loc['shape'])
        if loc['kind']=='cell':actual=shape.table.cell(loc['row'],loc['col']).text;cells+=1
        else:actual=shape.text
        expected=data['units'][loc['unit']]['text'][loc['start']:loc['end']]
        if clean(actual)!=clean(expected):errors.append((loc,expected,actual))
        covered[loc['unit']].append((loc['start'],loc['end']))
    assert not errors,errors[:3]
    # Comparación independiente: todos los nodos visibles del documento Word.
    with ZipFile(SOURCE) as wz:
        wx=etree.fromstring(wz.read('word/document.xml'))
        raw=''.join(wx.xpath('//w:t/text()',namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}))
    assert clean(raw)==clean(''.join(u['text'] for u in data['units'].values())), 'Texto Word no extraído íntegramente'
    # Cobertura carácter a carácter, excluyendo únicamente espacios de maquetación.
    missing=[]
    for uid,u in data['units'].items():
        if not u['text'].strip():continue
        present=[False]*len(u['text'])
        for a,b in covered.get(uid,[]):
            for k in range(a,b):present[k]=True
        lost=''.join(c for k,c in enumerate(u['text']) if not c.isspace() and not present[k])
        if lost:missing.append((uid,lost,u['text']))
    assert not missing,missing[:5]
    orig=[b for b in data['blocks'] if b['kind']=='image' and not b.get('reaction')]
    with ZipFile(OUT) as z:
        assert z.testzip() is None
        hashes={sha(z.read(n)) for n in z.namelist() if n.startswith('ppt/media/')}
        assert all(b['sha256'] in hashes for b in orig)
        media=len(hashes)
    visuals=[v for v in deck.visuals if v['kind']=='native_chemical']
    assert sorted(v['reaction'] for v in visuals)==list(range(1,11))
    chapter={}
    for n in range(16):
        nums=[i+1 for i,s in enumerate(deck.scenes) if s['chapter']==n]
        if nums:chapter[str(n)]={'inicio':min(nums),'fin':max(nums),'diapositivas':len(nums)}
    report={
       'fuente':SOURCE.name,'sha256_fuente_word':data['sha256'],
       'diapositivas':len(prs.slides),'secciones_del_resumen':15,
       'palabras_del_texto_fuente':sum(len(u['text'].split()) for u in data['units'].values()),
       'unidades_textuales_fuente':len([u for u in data['units'].values() if u['text'].strip()]),
       'cobertura_de_parrafos_y_celdas':'100%; sin letras, cifras o símbolos omitidos',
       'minimo_tipografico_texto_editable_puntos':min(allsize),
       'tamano_texto_principal_puntos':22,'tamano_tablas_puntos':20,
       'tablas_del_resumen':len([b for b in data['blocks'] if b['kind']=='table']),
       'tablas_nativas_powerpoint_paginadas':native_tables,
       'imagenes_originales_conservadas_sin_alteracion':len(orig),
       'esquemas_quimicos_vectoriales_editables':len(visuals),
       'vistas_ampliadas_adicionales':len([o for s in deck.scenes for o in s['objects'] if o['kind']=='image'])-len(orig),
       'errores_geometricos_detectados':len(deck.bbox_errors),
       'observacion_imagenes':'El mínimo tipográfico se verifica en texto editable y esquemas nativos. El texto rasterizado del libro conserva las limitaciones de la fuente comprimida; se incluyen vistas ampliadas y explicaciones legibles.',
       'mapa_por_seccion':chapter,'powerpoint':OUT.name,'pdf':PDF.name,
    }
    REPORT.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return report


def main():
    OUT.parent.mkdir(parents=True,exist_ok=True);CACHE.mkdir(parents=True,exist_ok=True)
    data=extract();composer=Composer(data);deck=composer.build()
    deck.prs.save(OUT)
    (CACHE/'trazabilidad.json').write_text(json.dumps(deck.ledger,ensure_ascii=False,indent=2),encoding='utf-8')
    (CACHE/'escenas.json').write_text(json.dumps(deck.scenes,ensure_ascii=False,indent=2),encoding='utf-8')
    result=verify(deck,data)
    deck.pdf(PDF)
    print(json.dumps(result,ensure_ascii=False,indent=2))
    print('PPTX MB:',round(OUT.stat().st_size/1e6,2),'PDF MB:',round(PDF.stat().st_size/1e6,2))

if __name__=='__main__':main()
