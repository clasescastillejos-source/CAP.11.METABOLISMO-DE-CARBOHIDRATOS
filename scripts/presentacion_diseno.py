"""Diseño de diapositivas y control geométrico; texto editable de al menos 18 pt.

El grafo de escenas permite producir una copia PDF con la misma composición.
Arial en PowerPoint; DejaVu Sans (métricas conservadoras) en el PDF de consulta.
"""
from __future__ import annotations
from pathlib import Path
import math
from copy import deepcopy
from dataclasses import dataclass, field
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

SW, SH = 960.0, 540.0  # 13 1/3 × 7 1/2 pulgadas, 16:9
LEFT, RIGHT = 46.0, 914.0
WHITE='FFFFFF'; NAVY='173B50'; INK='243C4C'; TEAL='007F88'; LIGHT='EEF5F7'; GRAY='536A78'; LINE='CCDDE5'; MAGENTA='A42365'; GOLD='BD812A'
COLORS={0:TEAL,1:TEAL,2:TEAL,3:MAGENTA,4:MAGENTA,5:MAGENTA,6:TEAL,7:GOLD,8:TEAL,9:'2868A0',10:'6D4491',11:TEAL,12:TEAL,13:'2868A0',14:NAVY,15:GRAY}
FONTDIR=Path('/usr/share/fonts/truetype/dejavu')
pdfmetrics.registerFont(TTFont('DV',str(FONTDIR/'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('DVB',str(FONTDIR/'DejaVuSans-Bold.ttf')))


def clean(text):
    return ''.join(text.split())


def spans_text(spans):
    return ''.join(s['text'] for s in spans)


def rich_slice(spans, start=0, end=None):
    if end is None: end=len(spans_text(spans))
    out=[]; pos=0
    for s in spans:
        lim=pos+len(s['text']); a=max(pos,start); b=min(lim,end)
        if b>a:
            copy={k:v for k,v in s.items()}; copy['text']=s['text'][a-pos:b-pos]; out.append(copy)
        pos=lim
    return out


def rich_width(spans, size):
    return sum(pdfmetrics.stringWidth(s['text'],'DVB' if s.get('bold') else 'DV',size) for s in spans)


def plain(text,bold=False,color=None):
    s={'text':text,'bold':bold}
    if color: s['color']=color
    return [s]


def wrap(spans,width,size):
    """Cortes explícitos sin eliminar letras, cifras ni símbolos del original."""
    text=spans_text(spans); out=[]; pos=0
    while pos<len(text):
        if text[pos]=='\n':
            out.append({'start':pos,'end':pos,'spans':[]});pos+=1;continue
        stop=text.find('\n',pos)
        if stop<0:stop=len(text)
        lo=pos+1; hi=stop; best=pos
        while lo<=hi:
            mid=(lo+hi)//2
            if rich_width(rich_slice(spans,pos,mid),size)<=width-2:best=mid;lo=mid+1
            else:hi=mid-1
        if best==pos:best=min(pos+1,stop)
        end=best
        if best<stop:
            breaks=[k for k in range(pos,best) if text[k].isspace() or text[k] in '-/']
            if breaks:
                k=breaks[-1]
                end=k if text[k].isspace() else k+1
            if end<=pos:end=best
        while end>pos and text[end-1].isspace():end-=1
        out.append({'start':pos,'end':end,'spans':rich_slice(spans,pos,end)})
        pos=end
        while pos<len(text) and text[pos].isspace() and text[pos]!='\n':pos+=1
        if pos<len(text) and text[pos]=='\n':pos+=1
    return out or [{'start':0,'end':0,'spans':[]}]


def text_height(lines,size,leading=1.16):
    return len(lines)*size*leading+3


class Deck:
    def __init__(self):
        self.prs=Presentation();self.prs.slide_width=Pt(SW);self.prs.slide_height=Pt(SH)
        self.prs.core_properties.title='Metabolismo de los hidratos de carbono — presentación íntegra del resumen'
        self.prs.core_properties.subject='Texto íntegro, figuras, cuadros, estructuras, cofactores y ejercicios'
        self.prs.core_properties.author='Presentación didáctica del capítulo proporcionado'
        self.scenes=[];self.ledger=[];self.pending_links=[];self.targets={};self.visuals=[];self.current=None
        self.bbox_errors=[]

    def new(self,chapter=0,title=None,dark=False,chapter_title=None,title_source=None):
        slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(NAVY if dark else WHITE)
        scene={'chapter':chapter,'dark':dark,'objects':[],'title':title or ''};self.scenes.append(scene);self.current=(slide,scene)
        accent=COLORS.get(chapter,TEAL)
        self.rect(0,0,SW,5,accent)
        if not dark:
            self.text('METABOLISMO DE LOS HIDRATOS DE CARBONO',LEFT,15,700,23,18,TEAL,bold=True)
            self.text('BIOQUÍMICA',770,15,144,23,18,GRAY,align='right')
            self.line(LEFT,45,RIGHT,45,LINE,1)
        if title:
            size=32;ls=wrap(plain(title,True),RIGHT-LEFT,size)
            while len(ls)>3 and size>26:
                size-=2;ls=wrap(plain(title,True),RIGHT-LEFT,size)
            h=text_height(ls,size,1.08)
            self.text(title,LEFT,59,RIGHT-LEFT,h,size,WHITE if dark else NAVY,bold=True,lines=ls,leading=1.08,source=title_source)
            top=59+h+17
        else:top=130
        scene['body_top']=max(125,top);scene['body_bottom']=481
        # Espacio de navegación, no parte del texto fuente.
        self.line(LEFT,498,RIGHT,498, '33586D' if dark else LINE,1)
        self.text('Texto íntegro del resumen · Presentación editable',LEFT,509,720,24,18,'D2E7ED' if dark else GRAY)
        return slide,scene

    def current_slide(self):return self.current[0]
    def current_scene(self):return self.current[1]

    def _bounds(self,x,y,w,h,label=''):
        if x<-.5 or y<-.5 or x+w>SW+.5 or y+h>SH+.5:
            self.bbox_errors.append((len(self.scenes),label,x,y,w,h))

    def rect(self,x,y,w,h,fill,line=None,radius=False):
        self._bounds(x,y,w,h,'rect')
        kind=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
        sh=self.current_slide().shapes.add_shape(kind,Pt(x),Pt(y),Pt(w),Pt(h))
        sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
        if line:sh.line.color.rgb=RGBColor.from_string(line);sh.line.width=Pt(.7)
        else:sh.line.fill.background()
        if radius:
            try:sh.adjustments[0]=.08
            except Exception:pass
        self.current_scene()['objects'].append({'kind':'rect','x':x,'y':y,'w':w,'h':h,'fill':fill,'line':line,'radius':radius})
        return sh

    def circle(self,x,y,d,fill,line=None):
        self._bounds(x,y,d,d,'circle')
        sh=self.current_slide().shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL,Pt(x),Pt(y),Pt(d),Pt(d))
        sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
        if line:sh.line.color.rgb=RGBColor.from_string(line)
        else:sh.line.fill.background()
        self.current_scene()['objects'].append({'kind':'circle','x':x,'y':y,'d':d,'fill':fill,'line':line})
        return sh

    def line(self,x1,y1,x2,y2,color=INK,width=1.6,arrow=False):
        sh=self.current_slide().shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2))
        sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
        if arrow:
            ln=sh._element.spPr.find(qn('a:ln'))
            el=OxmlElement('a:tailEnd');el.set('type','triangle');el.set('w','med');el.set('len','med');ln.append(el)
        self.current_scene()['objects'].append({'kind':'line','x1':x1,'y1':y1,'x2':x2,'y2':y2,'color':color,'width':width,'arrow':arrow})
        return sh

    def text(self,text,x,y,w,h,size=22,color=INK,bold=False,spans=None,lines=None,align='left',source=None,leading=1.16,name=None):
        assert size>=18,size
        self._bounds(x,y,w,h,'text '+text[:60])
        spans=spans if spans is not None else plain(text,bold)
        lines=lines if lines is not None else wrap(spans,w,size)
        assert max((rich_width(l['spans'],size) for l in lines),default=0)<=w+.1,(text,w,size)
        req=text_height(lines,size,leading)
        if req>h+2:self.bbox_errors.append((len(self.scenes),'texto supera altura: '+text[:60],req,h))
        sh=self.current_slide().shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h))
        tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Pt(0)
        tf.vertical_anchor=MSO_ANCHOR.TOP
        p=tf.paragraphs[0];p.font.name='Arial';p.font.size=Pt(size)
        p.line_spacing=Pt(size*leading);p.space_before=Pt(0);p.space_after=Pt(0)
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        for idx,l in enumerate(lines):
            if idx:p.add_line_break()
            for s in l['spans']:
                r=p.add_run();r.text=s['text'];r.font.name='Arial';r.font.size=Pt(size)
                r.font.bold=s.get('bold',False)
                r._r.get_or_add_rPr().set('lang','es-MX')
                r.font.color.rgb=RGBColor.from_string(s.get('color') or color)
                if s.get('link'):
                    r.hyperlink.address=s['link'];r.font.color.rgb=RGBColor.from_string('34D9CA' if self.current_scene()['dark'] else TEAL);r.font.underline=True
                if s.get('anchor'):self.pending_links.append((sh,s['anchor']))
        if source:
            uid,start,end=source;sh.name=f'FUENTE__{uid}__{start}__{end}'
            self.ledger.append({'unit':uid,'start':start,'end':end,'slide':len(self.scenes),'kind':'shape','shape':sh.name})
        elif name:sh.name=name
        obj={'kind':'text','text':text,'x':x,'y':y,'w':w,'h':h,'size':size,'color':color,'bold':bold,'spans':spans,'lines':lines,'leading':leading,'align':align,'source':source}
        self.current_scene()['objects'].append(obj)
        return sh

    def image(self,path,x,y,w,h,alt='',source_id=None):
        im=Image.open(path);iw,ih=im.size
        ratio=min(w/iw,h/ih);aw=iw*ratio;ah=ih*ratio
        xx=x+(w-aw)/2;yy=y+(h-ah)/2
        self._bounds(xx,yy,aw,ah,'image')
        pic=self.current_slide().shapes.add_picture(str(path),Pt(xx),Pt(yy),width=Pt(aw),height=Pt(ah))
        pic._element.nvPicPr.cNvPr.set('descr',alt)
        if source_id:pic.name='IMAGEN_FUENTE__'+source_id
        self.current_scene()['objects'].append({'kind':'image','path':str(path),'x':xx,'y':yy,'w':aw,'h':ah,'alt':alt,'source_id':source_id})
        return pic

    def table(self,rows,widths,x,y,heights,size=20,table_id='T',part=1):
        assert size>=18
        w=sum(widths);h=sum(heights);self._bounds(x,y,w,h,'table')
        sh=self.current_slide().shapes.add_table(len(rows),len(widths),Pt(x),Pt(y),Pt(w),Pt(h))
        sh.name=f'TABLA__{table_id}__{part}'
        t=sh.table;t.first_row=True;t.horz_banding=False
        for col,width in zip(t.columns,widths):col.width=Pt(width)
        for row,height in zip(t.rows,heights):row.height=Pt(height)
        source_map=[];draw_rows=[]
        for ri,row in enumerate(rows):
            dr=[]
            for ci,cell_info in enumerate(row):
                cell=t.cell(ri,ci);cell.margin_left=cell.margin_right=Pt(10);cell.margin_top=cell.margin_bottom=Pt(7)
                cell.vertical_anchor=MSO_ANCHOR.TOP
                header=ri==0
                bg=NAVY if header else (LIGHT if ri%2==0 else 'F8FAFB')
                if ci==0 and not header:bg='E8F2F4'
                cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(bg)
                cell.text_frame.clear();cell.text_frame.word_wrap=False;cell.text_frame.auto_size=MSO_AUTO_SIZE.NONE
                p=cell.text_frame.paragraphs[0];p.font.name='Arial';p.font.size=Pt(size);p.line_spacing=Pt(size*1.14)
                p.space_before=p.space_after=Pt(0)
                ss=cell_info['spans']
                if header:ss=[{**s,'bold':True,'color':WHITE} for s in ss]
                elif ci==0:ss=[{**s,'bold':True,'color':TEAL} for s in ss]
                ls=wrap(ss,widths[ci]-20,size)
                if text_height(ls,size,1.14)+14>heights[ri]+.5:
                    self.bbox_errors.append((len(self.scenes),'cell',table_id,ri,ci))
                for li,l in enumerate(ls):
                    if li:p.add_line_break()
                    for s in l['spans']:
                        r=p.add_run();r.text=s['text'];r.font.name='Arial';r.font.size=Pt(size);r.font.bold=s.get('bold',False)
                        r._r.get_or_add_rPr().set('lang','es-MX')
                        r.font.color.rgb=RGBColor.from_string(s.get('color') or (WHITE if header else INK))
                        if s.get('link'):r.hyperlink.address=s['link']
                if cell_info.get('source'):
                    uid,a,b=cell_info['source']
                    self.ledger.append({'unit':uid,'start':a,'end':b,'slide':len(self.scenes),'kind':'cell','shape':sh.name,'row':ri,'col':ci})
                dr.append({'spans':ss,'lines':ls,'text':spans_text(ss),'bg':bg})
            draw_rows.append(dr)
        self.current_scene()['objects'].append({'kind':'table','x':x,'y':y,'w':w,'h':h,'widths':widths,'heights':heights,'rows':draw_rows,'size':size,'leading':1.14})
        return sh

    def notes(self,text):
        frame=self.current_slide().notes_slide.notes_text_frame
        frame.text=text
        for p in frame.paragraphs:
            p.font.name='Arial';p.font.size=Pt(18)
            for r in p.runs:r.font.name='Arial';r.font.size=Pt(18)

    def finish(self):
        total=len(self.scenes)
        for i,(slide,scene) in enumerate(zip(self.prs.slides,self.scenes),1):
            self.current=(slide,scene)
            self.text(f'{i:03d} / {total:03d}',760,509,154,24,18,'D2E7ED' if scene['dark'] else GRAY,align='right')
        for shape,anchor in self.pending_links:
            if anchor in self.targets:shape.click_action.target_slide=self.targets[anchor]
        assert not self.bbox_errors,self.bbox_errors[:12]

    def pdf(self,path):
        c=canvas.Canvas(str(path),pagesize=(SW,SH),pageCompression=1)
        c.setTitle('Metabolismo de los hidratos de carbono — presentación íntegra')
        c.setAuthor('Material didáctico basado en el resumen de Word proporcionado')
        anchors={'sec_'+str(scene['chapter']):'slide'+str(i) for i,scene in enumerate(self.scenes,1) if scene.get('chapter_divider')}
        def draw_text(o):
            y=o['y']+o['size']*.89
            for l in o['lines']:
                width=rich_width(l['spans'],o['size'])
                x=o['x'] if o['align']=='left' else o['x']+(o['w']-width)*(1 if o['align']=='right' else .5)
                for s in l['spans']:
                    font='DVB' if s.get('bold') else 'DV';sz=o['size'];col=s.get('color') or o['color']
                    c.setFillColor(HexColor('#'+col));c.setFont(font,sz);c.drawString(x,SH-y,s['text'])
                    ww=pdfmetrics.stringWidth(s['text'],font,sz)
                    if s.get('link'):
                        c.linkURL(s['link'],(x,SH-y-3,x+ww,SH-y+sz),relative=0,thickness=0)
                    if s.get('anchor') in anchors:
                        c.linkRect('',anchors[s['anchor']],(x,SH-y-3,x+ww,SH-y+sz),relative=0,thickness=0)
                    x+=ww
                y+=o['size']*o['leading']
        for i,scene in enumerate(self.scenes,1):
            c.setFillColor(HexColor('#'+(NAVY if scene['dark'] else WHITE)));c.rect(0,0,SW,SH,fill=1,stroke=0)
            c.bookmarkPage(f'slide{i}')
            if scene.get('chapter_divider'):c.addOutlineEntry(scene['title'],f'slide{i}',level=0,closed=False)
            for o in scene['objects']:
                if o['kind']=='rect':
                    c.setFillColor(HexColor('#'+o['fill']))
                    if o['line']:c.setStrokeColor(HexColor('#'+o['line']));c.setLineWidth(.7)
                    if o['radius']:c.roundRect(o['x'],SH-o['y']-o['h'],o['w'],o['h'],8,fill=1,stroke=bool(o['line']))
                    else:c.rect(o['x'],SH-o['y']-o['h'],o['w'],o['h'],fill=1,stroke=bool(o['line']))
                elif o['kind']=='circle':
                    c.setFillColor(HexColor('#'+o['fill']))
                    if o['line']:c.setStrokeColor(HexColor('#'+o['line']))
                    c.circle(o['x']+o['d']/2,SH-o['y']-o['d']/2,o['d']/2,fill=1,stroke=bool(o['line']))
                elif o['kind']=='line':
                    c.setStrokeColor(HexColor('#'+o['color']));c.setLineWidth(o['width'])
                    c.line(o['x1'],SH-o['y1'],o['x2'],SH-o['y2'])
                    if o['arrow']:
                        dx=o['x2']-o['x1'];dy=o['y2']-o['y1'];ang=math.atan2(dy,dx);x,y=o['x2'],o['y2'];n=9
                        p=c.beginPath();p.moveTo(x,SH-y)
                        for a in [ang+2.65,ang-2.65]:p.lineTo(x+n*math.cos(a),SH-(y+n*math.sin(a)))
                        p.close();c.setFillColor(HexColor('#'+o['color']));c.drawPath(p,fill=1,stroke=0)
                elif o['kind']=='image':c.drawImage(ImageReader(o['path']),o['x'],SH-o['y']-o['h'],o['w'],o['h'],preserveAspectRatio=True,mask='auto')
                elif o['kind']=='text':draw_text(o)
                elif o['kind']=='table':
                    yy=o['y']
                    for row,height in zip(o['rows'],o['heights']):
                        xx=o['x']
                        for cell,width in zip(row,o['widths']):
                            c.setFillColor(HexColor('#'+cell['bg']));c.setStrokeColor(HexColor('#'+LINE));c.setLineWidth(.55)
                            c.rect(xx,SH-yy-height,width,height,fill=1,stroke=1)
                            draw_text({'x':xx+10,'y':yy+7,'w':width-20,'h':height-14,'size':o['size'],'leading':o['leading'],'align':'left','color':INK,'lines':cell['lines']})
                            xx+=width
                        yy+=height
            c.showPage()
        c.save()
